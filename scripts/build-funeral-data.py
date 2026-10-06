"""전국 장례식장·장지 목록 → data/funeral-homes.json, data/burial-places.json

원본: 공공데이터포털 '재단법인한국장례문화진흥원_전국 장사시설 현황 시설_가격정보_20230602'
      (https://www.data.go.kr/data/15021763/fileData.do, 공공저작물 출처표시 1유형)
      압축 안의 '1.장사시설 현황_20230601.xlsx'.
사용: python3 scripts/build-funeral-data.py <장사시설 현황.xlsx>
card-make.html 의 장례식장·장지 검색이 이 파일을 읽는다(키·서버 없이 자동완성).
"""
import json, re, sys
import openpyxl

SRC = "재단법인한국장례문화진흥원 전국 장사시설 현황(2023-06-01), 공공데이터포털"
PREFIX = re.compile(r"^\s*\((?:유|주|재|사|복|의|사단|재단|학)\)\s*")

def clean(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()

def rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [clean(h) for h in next(it)]
    for r in it:
        d = {head[i]: clean(v) for i, v in enumerate(r) if i < len(head)}
        if d.get("시설명"):
            yield d

def main(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    homes = []
    for d in rows(wb["장례식장"]):
        name = PREFIX.sub("", d["시설명"]).strip()
        homes.append({
            "n": name,
            "a": d.get("주소", ""),
            "t": d.get("전화번호", ""),
            "r": int(d["빈소수"]) if d.get("빈소수", "").isdigit() else None,
            "k": d.get("운영종류", ""),   # 병원 / 전문 등
        })
    homes.sort(key=lambda x: x["n"])
    burial = []
    for sheet, kind in (("봉안시설", "봉안"), ("묘지", "묘지"), ("자연장지", "자연장"), ("화장시설", "화장")):
        if sheet not in wb.sheetnames:
            continue
        for d in rows(wb[sheet]):
            burial.append({"n": PREFIX.sub("", d["시설명"]).strip(), "a": d.get("주소", ""), "k": kind})
    burial.sort(key=lambda x: x["n"])
    for fn, items in (("data/funeral-homes.json", homes), ("data/burial-places.json", burial)):
        with open(fn, "w", encoding="utf-8") as f:
            json.dump({"source": SRC, "count": len(items), "items": items}, f, ensure_ascii=False, separators=(",", ":"))
        print(fn, len(items))

if __name__ == "__main__":
    main(sys.argv[1])
